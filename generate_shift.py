#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
グラウンドアップ シフト表 自動生成スクリプト（v2）
横浜市 地域密着型通所介護「従業者の勤務の体制及び勤務形態一覧表」形式

【このスクリプトが自動化する範囲（v2）】
  ステップ1: リセット作業（月の切り替え、パートの固定曜日・常勤の出勤日を組み立てる）
  ステップ2: 休み希望・有休・勤務可能希望をそのまま反映する
  ステップ3: 常勤スタッフの所定勤務日数（有休を除く実質所定日数）に合わせて自動調整。
             候補日は「営業日」全体から、その時点で配置人数が最も少ない日を優先。
             6日以上の連続勤務にならないようガードする。
  ステップ4: 1日あたりの配置人数（Total）が4〜6人になるよう、
             スタッフが「勤務可能」欄に記載した日を優先的に組み込んで調整する
             （記載がない場合は自動では人を追加しない。4人未満の営業日は
             出力Excel上で赤く強調表示し、レポートにも警告を出す）。
  完成後、独立したレビュー関数で「所定日数」「Total人数」「連続勤務」
  「兼務者の勤務日一致」などをもう一度チェックし、レポートに結果を出す。

【営業日の定義（ユーザー確定、2026-09時点）】
  日曜、および下記の固定休業期間を除く日:
    5/3・5/4・5/5（GW）、8/13・8/14・8/15（お盆）、
    12/30・12/31（年末）、1/1・1/2・1/3（年始）
  ※ この「営業日」は「スタッフに勤務させて良い日」の判定に使う。
    常勤スタッフの所定勤務日数（下記）は、これとは別に「土日祝日
    （国民の祝日を含む）」の日数を使って計算する、ユーザー確定の別ルール。

【常勤スタッフの所定勤務日数】
  所定勤務日数 = その月の日数 - (土曜日数 + 日曜日数 + 祝日日数)
  例: 2026年10月は31日中、土日祝日が10日（土5・日4・祝1）なので所定21日・168時間。

【使い方】
  python3 generate_shift.py --base シフト2026.xlsx --source-am 9月1 --source-pm 9月2 \
      --year 2026 --month 10 --requests requests_202610.csv \
      --staff staff_master.csv --out output/シフト2026_10月.xlsx

  requests.csv の形式: name,date(YYYY-MM-DD),status
    status:
      off         … その日は休み
      paid_leave  … 有休（公休＝土日祝日とは別枠の休みとして扱う）
      am_only     … 午前のみ出勤可（午後シートからは外す）
      pm_only     … 午後のみ出勤可（午前シートからは外す）
      available   … 勤務可能（休み希望とは別の欄。記載がある日は自動的に
                    その人のシフトに組み込む。Total人数が手薄な日の
                    穴埋めに使う）
"""
import argparse
import calendar
import csv
import datetime
import sys
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.styles import PatternFill, Font
from openpyxl.formatting.rule import CellIsRule

try:
    import jpholiday
except ImportError:
    jpholiday = None

WEEKDAY_JP = ['月', '火', '水', '木', '金', '土', '日']  # Python weekday(): Mon=0
FIRST_DAY_COL = 19  # column S
STAFF_ROWS_FIRST = 25
STAFF_ROWS_END = 80  # この行(80)から「（参考）1日の職種別人員内訳」の集計行になるため、
                      # スタッフ行のスキャン・クリアはこれより前で必ず止めること。
                      # （行80=生活相談員, 81=看護職員, 82=介護職員, 83=機能訓練指導員,
                      #   84=Total はいずれもCOUNTIFS/SUM数式。誤って25〜83をスタッフ行と
                      #   みなしてクリアすると、行80の数式を壊してTotalが過小集計になる
                      #   重大なバグがあったため、明示的な境界を設ける。）
TOTAL_ROW = 84
# (15)利用者数（予定作成時は定員数）、(16)サービス提供時間（平均提供時間）の行。
# ベース様式では日付列ごとに固定値（例: 定員15人、平均提供時間3時間）が入っており、
# 日曜（および固定休業日）の列だけ空欄になっている。copy_worksheet()はこの値を
# そのままコピーするだけなので、複製元の月と対象月とで曜日がずれると、日曜以外の
# 列が空欄のまま/日曜の列に値が残ったままになってしまう。ユーザー確定ルール
# （2026-10）：対象月の日曜（営業日でない日）は0にする。
CAPACITY_ROW = 77
SERVICE_TIME_ROW = 78
MAX_CONSECUTIVE_WORK_DAYS = 5  # 「6日以上の連続勤務」を避けるため、連続は5日まで
MIN_DAILY_TOTAL = 4
MAX_DAILY_TOTAL = 6

# 有休（paid_leave）の日にシフト表へ書き込むシフト記号。shift_codes.csv側で
# 4.0時間（AM/PM共通）として定義済み（ユーザー確定ルール、2026-10）。
# これにより月合計・週平均勤務時間数（AX/AZ列）にも有休分の時間が反映され、
# 常勤スタッフの月間時間数が有休の有無に関わらず所定どおり（例: 168h/月の人は
# 週平均・月合計とも一律の値）になる。
PAID_LEAVE_CODE = '有'

# 事業所の固定休業日（年をまたいで毎年同じ月日）。ユーザー確定ルール（2026-09）。
FIXED_CLOSURE_MD = {
    (5, 3), (5, 4), (5, 5),        # ゴールデンウィーク
    (8, 13), (8, 14), (8, 15),     # お盆
    (12, 30), (12, 31),            # 年末
    (1, 1), (1, 2), (1, 3),        # 年始
}

RED_FILL = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
RED_FONT = Font(color='9C0006', bold=True)


# ---------------------------------------------------------------------------
# 基本ユーティリティ
# ---------------------------------------------------------------------------

def is_business_day(d):
    """営業日か（日曜・固定休業期間ではないか）を判定する"""
    if d.weekday() == 6:  # 日曜
        return False
    if (d.month, d.day) in FIXED_CLOSURE_MD:
        return False
    return True


def business_days_list(year, month):
    days_in_month = calendar.monthrange(year, month)[1]
    return [day for day in range(1, days_in_month + 1)
            if is_business_day(datetime.date(year, month, day))]


def full_time_target_days(year, month):
    """常勤スタッフの所定勤務日数を計算する。

    ルール（ユーザー確定、2026-10時点）:
      事業所は日曜以外営業。常勤は「土日祝日の日数分」休みを取る
      （土曜や祝日に出た分は別の平日で調整して帳尻を合わせる）。
      → 所定勤務日数 = その月の日数 - (土曜+日曜+祝日の日数)
      例: 2026年10月は31日中、土日祝日が10日なので所定21日。

    ※ ここでの「祝日」は国民の祝日（jpholidayで判定）全部を指し、上の
      「営業日」の固定休業カレンダーとは別の概念。事業所自体はその祝日も
      営業している場合があるが、常勤スタッフの休日クオータとしてカウントする。
    """
    days_in_month = calendar.monthrange(year, month)[1]
    off_days = 0
    for day in range(1, days_in_month + 1):
        d = datetime.date(year, month, day)
        wd = d.weekday()  # Mon=0 ... Sun=6
        is_weekend = wd in (5, 6)  # Sat, Sun
        is_holiday = jpholiday is not None and jpholiday.is_holiday(d)
        if is_weekend or is_holiday:
            off_days += 1
    return days_in_month - off_days, off_days


def max_consecutive_run(days_set):
    """連続勤務日数の最大値を返す"""
    if not days_set:
        return 0
    s = sorted(days_set)
    best = cur = 1
    for i in range(1, len(s)):
        if s[i] == s[i - 1] + 1:
            cur += 1
            best = max(best, cur)
        else:
            cur = 1
    return best


def load_service_hours_table(path='shift_codes.csv'):
    """シフト記号 → サービス提供時間内の勤務時間数（19列目のVLOOKUP対象）を読み込む。

    横浜市様式の「Total」（(18) 1日の職種別人員内訳の合計）は、シフト記号の
    有無ではなく、この「サービス提供時間内の勤務時間数」が0より大きいかどうかで
    COUNTIFS判定している。例えば管理者の事務作業用シフト記号 b（8:30〜9:00）は
    サービス提供時間（9:00〜）に掛かっていないため0.0時間となり、bで出勤していても
    Totalにはカウントされない。この関数はそれを再現するための対応表を返す。
    """
    table = {}
    try:
        with open(path, newline='', encoding='utf-8-sig') as f:
            for row in csv.DictReader(f):
                code = row.get('code', '').strip()
                if not code:
                    continue
                try:
                    hours = float(row.get('service_work_hours', '0') or 0)
                except ValueError:
                    hours = 0.0
                table[code] = hours
    except FileNotFoundError:
        pass
    return table


def counts_toward_total(code, role, service_hours_table, total_roles):
    """このシフト記号・職種が横浜市様式のTotal（配置人数、row84）にカウントされるか

    Total行は (a) L80:L83 に列挙された職種（生活相談員・看護職員・介護職員・
    機能訓練指導員）に属し、かつ (b) その日の記号のサービス提供時間内の勤務時間数が
    0より大きい、という2条件を両方満たす行だけをCOUNTIFSで数えている。
    管理者・送迎職員はそもそも(a)に該当しないため、コードに関わらず常に対象外。
    """
    if not code:
        return False
    if role not in total_roles:
        return False
    return service_hours_table.get(code, 1.0) > 0  # 表にない記号は安全側でカウントする


def load_staff_master(path):
    staff = []
    with open(path, newline='', encoding='utf-8-sig') as f:
        for row in csv.DictReader(f):
            staff.append(row)
    return staff


def load_requests(path):
    """name -> {date(iso): status}"""
    reqs = {}
    if not path:
        return reqs
    with open(path, newline='', encoding='utf-8-sig') as f:
        for row in csv.DictReader(f):
            reqs.setdefault(row['name'], {})[row['date']] = row['status']
    return reqs


# ---------------------------------------------------------------------------
# 休み希望・勤務可能希望
# ---------------------------------------------------------------------------

def _days_with_status(staff_name, requests, year, month, statuses):
    reqs = requests.get(staff_name, {})
    days = set()
    for date_str, status in reqs.items():
        try:
            d = datetime.date.fromisoformat(date_str)
        except ValueError:
            continue
        if d.year == year and d.month == month and status in statuses:
            days.add(d.day)
    return days


def requested_blocked_days(staff_name, requests, year, month, unit):
    """休み希望などで、その単位（AM/PM）に出勤できない日の集合を返す"""
    reqs = requests.get(staff_name, {})
    blocked = set()
    for date_str, status in reqs.items():
        try:
            d = datetime.date.fromisoformat(date_str)
        except ValueError:
            continue
        if d.year != year or d.month != month:
            continue
        day = d.day
        if status in ('off', 'paid_leave'):
            blocked.add(day)
        elif status == 'am_only' and unit == 'PM':
            blocked.add(day)
        elif status == 'pm_only' and unit == 'AM':
            blocked.add(day)
    return blocked


def paid_leave_days(staff_name, requests, year, month):
    """有休（paid_leave）とマークされた日の集合を返す。

    有休は「公休（土日祝日）」とは別枠の休みなので、所定勤務日数からその分を
    差し引いた「実質の所定日数」と比較する（ユーザー確定ルール、2026-10）。
    """
    return _days_with_status(staff_name, requests, year, month, {'paid_leave'})


def available_request_days(staff_name, requests, year, month):
    """「勤務可能」欄に記載された日の集合を返す（休み希望とは別枠）。

    記載がある日は、営業日である限りその人のシフトに自動的に組み込む
    （Total人数が手薄な日の穴埋めに使う、ユーザー確定ルール、2026-09）。
    """
    return _days_with_status(staff_name, requests, year, month, {'available'})


# ---------------------------------------------------------------------------
# シート操作
# ---------------------------------------------------------------------------

def clear_and_write_headers(ws, year, month):
    ws['AC2'] = year
    ws['AG2'] = month

    # (3)事業所における常勤の従業者が勤務すべき時間数（BB6, 時間/月）は複製元シートの
    # 値をそのまま引き継いでしまう静的な数値セルなので、対象月の所定勤務日数に
    # 合わせて毎回上書きする（ユーザー確定ルール、2026-10: 所定勤務日数×8時間）。
    # 例: 2026年10月は所定21日 → 21×8=168時間。
    target_days, _ = full_time_target_days(year, month)
    ws['BB6'] = target_days * 8

    days_in_month = calendar.monthrange(year, month)[1]

    for day in range(1, 32):
        col = get_column_letter(FIRST_DAY_COL + day - 1)
        if day > days_in_month:
            ws[f'{col}22'] = None
            ws[f'{col}23'] = None
            ws[f'{col}24'] = None
            continue
        d = datetime.date(year, month, day)
        ws[f'{col}22'] = day
        ws[f'{col}23'] = (d.weekday() % 7) + 1
        ws[f'{col}24'] = WEEKDAY_JP[d.weekday()]
    return days_in_month


def clear_staff_shift_rows(ws, max_row=STAFF_ROWS_END):
    """各スタッフのシフト記号行（3行1組の1行目）のS〜AW列を空にし、
    Totalの赤強調も一旦クリアする。

    注意: B列（No）は1行目だけ固定値で、2行目以降は "=B25+1" のような数式が
    入っているため、B列がintかどうかでは判定できない。氏名（L列）の有無で判定する。
    max_row は必ずSTAFF_ROWS_END（80）以下にすること。80行目以降は
    「（参考）1日の職種別人員内訳」の集計行（COUNTIFS/SUM数式）で、
    スタッフ行ではない。
    """
    r = STAFF_ROWS_FIRST
    while r < max_row:
        name = ws[f'L{r}'].value
        if name:
            for col_idx in range(FIRST_DAY_COL, FIRST_DAY_COL + 31):
                col = get_column_letter(col_idx)
                ws[f'{col}{r}'] = None
            r += 3
        else:
            r += 1
    for col_idx in range(FIRST_DAY_COL, FIRST_DAY_COL + 31):
        col = get_column_letter(col_idx)
        cell = ws[f'{col}{TOTAL_ROW}']
        cell.fill = PatternFill(fill_type=None)
        cell.font = Font()


def rewrite_capacity_and_service_time(ws, days_in_month, biz_days):
    """(15)利用者数・(16)サービス提供時間の行を対象月の曜日に合わせて書き直す。

    複製元シート（例: 9月分）の値をそのままコピーすると、非営業日（日曜・固定休業日）
    の列がコピー元の月の曜日ズレのまま残ってしまう（複製元で空欄だった列が対象月では
    営業日だったり、その逆だったりする）。ここで複製元シートに実際に入っている
    「営業日の値」（定員数・平均提供時間、通常は月をまたいで一定）を読み取り、
    対象月の営業日はその値、非営業日（日曜・固定休業日）は0で書き直す
    （ユーザー確定ルール、2026-10）。月の日数を超える列（29〜31日目）は空欄にする。
    """
    for row in (CAPACITY_ROW, SERVICE_TIME_ROW):
        # 複製元シートに実際に入っている「営業日の値」を読み取る（通常は単一の定数）。
        values = []
        for col_idx in range(FIRST_DAY_COL, FIRST_DAY_COL + 31):
            col = get_column_letter(col_idx)
            v = ws[f'{col}{row}'].value
            if v is not None:
                values.append(v)
        if not values:
            continue
        # 複製元シートで最も多く使われていた値を「営業日の値」として採用する
        biz_value = max(set(values), key=values.count)
        for day in range(1, 32):
            col = get_column_letter(FIRST_DAY_COL + day - 1)
            if day > days_in_month:
                ws[f'{col}{row}'] = None
            elif day in biz_days:
                ws[f'{col}{row}'] = biz_value
            else:
                ws[f'{col}{row}'] = 0


def default_pattern_for_part_time(staff_row, unit, year, month, biz_days):
    """パートスタッフの基本パターン（固定曜日）を営業日ベースで展開する"""
    code_col = 'default_code_am' if unit == 'AM' else 'default_code_pm'
    code = staff_row.get(code_col, '').strip()
    if not code:
        return {}
    fixed_col = 'fixed_weekdays_am' if unit == 'AM' else 'fixed_weekdays_pm'
    fixed_weekdays = [w.strip() for w in staff_row.get(fixed_col, '').split(',') if w.strip()]
    if not fixed_weekdays:
        return {}
    result = {}
    for day in biz_days:
        d = datetime.date(year, month, day)
        wd = WEEKDAY_JP[d.weekday()]
        if wd in fixed_weekdays:
            result[day] = code
    return result


def write_pattern(ws, row, pattern):
    for day, code in pattern.items():
        col = get_column_letter(FIRST_DAY_COL + day - 1)
        ws[f'{col}{row}'] = code


def highlight_out_of_range(ws, daily_counts, biz_days):
    """Total（row84）が4〜6人の範囲外（4人未満 または 7人以上）の営業日を
    赤く強調表示する。戻り値: (understaffed_days, overstaffed_days)"""
    understaffed = []
    overstaffed = []
    for day in biz_days:
        count = daily_counts.get(day, 0)
        if count < MIN_DAILY_TOTAL or count > MAX_DAILY_TOTAL:
            col = get_column_letter(FIRST_DAY_COL + day - 1)
            cell = ws[f'{col}{TOTAL_ROW}']
            cell.fill = RED_FILL
            cell.font = RED_FONT
            if count < MIN_DAILY_TOTAL:
                understaffed.append(day)
            else:
                overstaffed.append(day)
    return understaffed, overstaffed


def add_total_conditional_formatting(ws, days_in_month):
    """Total（row84）に、生成時点の状態に関わらず常に正しく判定されるよう、
    Excelの条件付き書式（4人未満・7人以上を赤く強調）を追加する。

    highlight_out_of_range() は生成した時点のdaily_countsに基づいて一度だけ
    セルを赤く塗るだけなので、生成後にExcel上で手作業によりシフト記号を
    書き換えた場合（有休の追加・変更を含む）は追従しない。Total自体は
    様式のCOUNTIFS数式で常に再計算される（有休の記号「有」もサービス提供時間内の
    勤務時間数が0より大きいためカウントされる）ため、条件付き書式を使えば
    手作業での修正後もExcel上で自動的に赤字判定が更新される
    （ユーザー確定ルール、2026-10：「有給を含んだ人数でTotalが4未満になる場合も
    赤く表示する」を、生成時だけでなく常に成り立つようにするため）。
    """
    if days_in_month <= 0:
        return
    first_col = get_column_letter(FIRST_DAY_COL)
    last_col = get_column_letter(FIRST_DAY_COL + days_in_month - 1)
    cell_range = f'{first_col}{TOTAL_ROW}:{last_col}{TOTAL_ROW}'
    ws.conditional_formatting.add(
        cell_range,
        CellIsRule(operator='lessThan', formula=[str(MIN_DAILY_TOTAL)], fill=RED_FILL, font=RED_FONT),
    )
    ws.conditional_formatting.add(
        cell_range,
        CellIsRule(operator='greaterThan', formula=[str(MAX_DAILY_TOTAL)], fill=RED_FILL, font=RED_FONT),
    )


# ---------------------------------------------------------------------------
# 常勤スタッフの自動配置（ステップ3）
# ---------------------------------------------------------------------------

def _pick_best_day(group, biz_days, daily_counts, allow_run_violation=False):
    candidates = [d for d in biz_days
                  if d not in group['blocked'] and d not in group['days']]
    candidates.sort(key=lambda d: (daily_counts[d], d))
    for d in candidates:
        trial = group['days'] | {d}
        if max_consecutive_run(trial) <= MAX_CONSECUTIVE_WORK_DAYS:
            return d
    if allow_run_violation and candidates:
        return candidates[0]
    return None


def allocate_full_time_groups(groups, biz_days, daily_counts):
    """常勤スタッフ（兼務者は氏名単位でグループ化済み）を、
    実質所定日数まで、その時点で最も手薄な営業日から順にラウンドロビンで埋める。
    6日以上連続勤務にならないよう候補日をガードする。
    """
    for g in groups:
        g.setdefault('run_violation', False)

    active = [g for g in groups if len(g['days']) < g['target']]
    while active:
        progressed = False
        for g in active:
            if len(g['days']) >= g['target']:
                continue
            day = _pick_best_day(g, biz_days, daily_counts)
            if day is None:
                # 連続勤務ガードで詰まった場合のみ、緩めて確保する（要確認フラグを立てる）
                day = _pick_best_day(g, biz_days, daily_counts, allow_run_violation=True)
                if day is not None:
                    g['run_violation'] = True
            if day is not None:
                g['days'].add(day)
                daily_counts[day] = daily_counts.get(day, 0) + g['weight']
                progressed = True
        active = [g for g in active if len(g['days']) < g['target']]
        if not progressed:
            break
    for g in groups:
        g['shortage'] = max(0, g['target'] - len(g['days']))
    return groups


# ---------------------------------------------------------------------------
# 月単位（AM または PM）の処理
# ---------------------------------------------------------------------------

def process_unit(wb, source_sheet_name, new_sheet_name, unit, staff_master, requests, year, month,
                  svc_table):
    src = wb[source_sheet_name]
    if new_sheet_name in wb.sheetnames:
        del wb[new_sheet_name]
    ws = wb.copy_worksheet(src)
    ws.title = new_sheet_name

    # --- コピー元シート名へのハードコードされたシート参照を自己参照に修正 ---
    # ベース様式（シフト2026.xlsx）の「(12)週平均勤務時間数」欄（AZ列、暦月モード）
    # など一部の数式は、本来AX列と同様に自シート内で完結すべきところ、
    # テンプレート作成時の入力ミスと思われる形で複製元シート名（'9月1'!や'9月2'!）
    # を直接埋め込んでしまっている。copy_worksheet()は数式を文字列としてそのまま
    # コピーするため、毎月このシート名を書き換えないと、複製元シートが出力に
    # 含まれなくなった時点で#NAME?エラーになる（2026-10、ユーザー報告で発覚）。
    # ここでシート参照部分（'9月1'!など）を取り除き、自シート内の自己参照に直す。
    old_sheet_ref = f"'{source_sheet_name}'!"
    for row in ws.iter_rows():
        for cell in row:
            v = cell.value
            if isinstance(v, str) and v.startswith('=') and old_sheet_ref in v:
                cell.value = v.replace(old_sheet_ref, '')

    days_in_month = clear_and_write_headers(ws, year, month)
    clear_staff_shift_rows(ws)

    biz_days = business_days_list(year, month)
    target_days, _ = full_time_target_days(year, month)

    # (15)利用者数・(16)サービス提供時間を対象月の曜日に合わせて書き直す
    # （日曜・固定休業日は0にする。ユーザー確定ルール、2026-10）。
    rewrite_capacity_and_service_time(ws, days_in_month, biz_days)

    # 「Total」（(18) 1日の職種別人員内訳の合計、row84）は L80:L83 に列挙された
    # 職種（生活相談員・看護職員・介護職員・機能訓練指導員）のみを合計しており、
    # 管理者や送迎職員はそもそも集計対象に入っていない（様式の構造）。
    # これをシートから毎回読み取ることで、様式が変わっても追従できるようにする。
    total_roles = {ws[f'L{r}'].value for r in range(80, TOTAL_ROW) if ws[f'L{r}'].value}

    # --- 氏名の出現回数（同名複数行の誤マッチを防ぐため） ---
    # STAFF_ROWS_END（80行目）以降は「（参考）1日の職種別人員内訳」の集計行
    # （COUNTIFS/SUM数式）なので、スタッフ行として絶対にスキャンしないこと。
    name_counts = {}
    r = STAFF_ROWS_FIRST
    while r < STAFF_ROWS_END:
        nm = ws[f'L{r}'].value
        if nm:
            name_counts[nm] = name_counts.get(nm, 0) + 1
            r += 3
        else:
            r += 1

    # --- シート行 ↔ staff_master のマッチング ---
    rows_info = []
    unmatched = []
    r = STAFF_ROWS_FIRST
    while r < STAFF_ROWS_END:
        name = ws[f'L{r}'].value
        if name:
            role = ws[f'C{r}'].value
            # 氏名+職種の完全一致を優先。氏名がシート内で1回しか出てこない場合のみ
            # 氏名だけでのフォールバックを許可する（同名複数行での誤マッチ事故防止）。
            staff_row = next((s for s in staff_master if s['name'] == name and s.get('role') == role), None)
            if staff_row is None and name_counts.get(name, 0) == 1:
                staff_row = next((s for s in staff_master if s['name'] == name), None)
            if staff_row is None:
                unmatched.append((r, name, role))
                r += 3
                continue
            rows_info.append({
                'row': r,
                'name': name,
                'role': role,
                'staff_row': staff_row,
                'is_full_time': staff_row.get('keitai') in ('A', 'B'),
            })
            r += 3
        else:
            r += 1

    # --- 退職者など、staff_master.csvに存在しない氏名の行を出力シートから消す ---
    # 「他の行で正しくマッチしている（＝現職スタッフの未使用重複スロット）」場合は
    # 触らない。それ以外（＝どの行にも一致しない、本当に退職者などの可能性が高い行）
    # だけ、氏名・職種・勤務形態記号・資格の欄を空にする。複製元（実績データが
    # 入ったベースファイル）は一切変更しない（ws は複製先の新シートのみ）。
    # ユーザー確定ルール（2026-10）：退職者の氏名が新しい月に毎回コピーされ続ける
    # 問題を、複製元の実績データを消さずに解消するための措置。
    matched_names = {info['name'] for info in rows_info}
    cleared_rows = []
    for r, name, role in unmatched:
        if name not in matched_names:
            for col in ('C', 'G', 'H', 'L'):
                ws[f'{col}{r}'] = None
            cleared_rows.append((r, name, role))
    # 自動で消した行は、以降のレビュー・レポートでは「未対応の欠落」として
    # 二重に警告しない（すでに解消済みのため）。
    unmatched = [u for u in unmatched if u not in cleared_rows]

    code_col = 'default_code_am' if unit == 'AM' else 'default_code_pm'

    # --- パートスタッフ: 固定曜日 + 勤務可能希望 − 休み希望 ---
    # 「Total」（配置人数）は、横浜市様式の数式と同じく「サービス提供時間内の
    # 勤務時間数」が0より大きいシフト記号のときだけカウントする（例: 管理者の
    # 事務作業用記号bは0.5h勤務でもサービス提供時間にかかっていないため0扱い）。
    daily_counts = {d: 0 for d in range(1, days_in_month + 1)}
    for info in rows_info:
        if info['is_full_time']:
            continue
        blocked = requested_blocked_days(info['name'], requests, year, month, unit)
        pattern = default_pattern_for_part_time(info['staff_row'], unit, year, month, biz_days)
        code = info['staff_row'].get(code_col, '').strip()
        avail = available_request_days(info['name'], requests, year, month)
        for day in avail:
            if day in biz_days and day not in blocked:
                pattern.setdefault(day, code)
        for day in blocked:
            pattern.pop(day, None)
        # 有休の日もシフト表に「有」と明記し、月合計・週平均勤務時間数（AX/AZ列）に
        # 反映されるようにする（パートスタッフも常勤と同様の扱い、ユーザー確定ルール、2026-10）。
        pl_days = paid_leave_days(info['name'], requests, year, month)
        for day in pl_days:
            if day in biz_days:
                pattern[day] = PAID_LEAVE_CODE
        info['pattern'] = pattern
        info['blocked'] = blocked
        info['code'] = code
        row_weight = 1 if counts_toward_total(code, info['role'], svc_table, total_roles) else 0
        pl_weight = 1 if counts_toward_total(PAID_LEAVE_CODE, info['role'], svc_table, total_roles) else 0
        for day, day_code in pattern.items():
            daily_counts[day] += pl_weight if day_code == PAID_LEAVE_CODE else row_weight

    # --- 常勤スタッフ: 氏名単位でグループ化（兼務者は勤務日を必ず揃える） ---
    ft_groups_by_name = {}
    for info in rows_info:
        if not info['is_full_time']:
            continue
        ft_groups_by_name.setdefault(info['name'], []).append(info)

    groups = []
    for name, infos in ft_groups_by_name.items():
        blocked = requested_blocked_days(name, requests, year, month, unit)
        pl_days = paid_leave_days(name, requests, year, month)
        target_effective = max(0, target_days - len(pl_days))
        avail = available_request_days(name, requests, year, month)
        avail = {d for d in avail if d in biz_days and d not in blocked}
        for info in infos:
            info['code'] = info['staff_row'].get(code_col, '').strip()
        # グループ全体としてのTotal寄与ウェイト＝行ごとの記号がサービス提供時間に
        # 掛かっている行の数（例: 上口さんは生活相談員行(c)のみカウントし、
        # 管理者行(b)はカウントしない → weight=1）
        weight = sum(1 for info in infos if counts_toward_total(info['code'], info['role'], svc_table, total_roles))
        # 有休（PAID_LEAVE_CODE）を書き込む日のTotal寄与ウェイトも別途計算しておく
        # （有休の記号はshift_codes.csv上サービス提供時間内の勤務時間数が0より
        # 大きいため、Totalにカウントされる。ユーザー確定ルール、2026-10）。
        pl_weight = sum(1 for info in infos
                         if counts_toward_total(PAID_LEAVE_CODE, info['role'], svc_table, total_roles))
        group = {
            'name': name,
            'infos': infos,
            'weight': weight,
            'pl_weight': pl_weight,
            'blocked': blocked,
            'target': target_effective,
            'pl_days': pl_days,
            'days': set(avail),
        }
        groups.append(group)
        for day in avail:
            daily_counts[day] += group['weight']
        for day in pl_days:
            if day in biz_days:
                daily_counts[day] += pl_weight

    groups.sort(key=lambda g: g['name'])
    allocate_full_time_groups(groups, biz_days, daily_counts)

    for g in groups:
        for info in g['infos']:
            code = info['staff_row'].get(code_col, '').strip()
            info['pattern'] = {day: code for day in sorted(g['days'])}
            # 有休の日もシフト表に「有」と明記し、月合計・週平均勤務時間数（AX/AZ列）に
            # 反映されるようにする（ユーザー確定ルール、2026-10）。
            for day in sorted(g['pl_days']):
                if day in biz_days:
                    info['pattern'][day] = PAID_LEAVE_CODE
            info['blocked'] = g['blocked']

    # --- Total（配置人数）が4〜6人の範囲外（4人未満・7人以上）の営業日を赤強調 ---
    for info in rows_info:
        write_pattern(ws, info['row'], info['pattern'])
    understaffed_days, overstaffed_days = highlight_out_of_range(ws, daily_counts, biz_days)
    # 生成後にExcel上で手作業修正しても赤字判定が追従するよう、条件付き書式も設定する
    add_total_conditional_formatting(ws, days_in_month)

    all_patterns = {}
    code_by_key = {}
    role_by_key = {}
    shortages = {}
    run_violations = {}
    for info in rows_info:
        key = f"{info['name']}（{info['role']}）" if info['role'] else info['name']
        all_patterns[key] = info['pattern']
        code_by_key[key] = info['code']
        role_by_key[key] = info['role']
    for g in groups:
        if g.get('shortage'):
            for info in g['infos']:
                key = f"{info['name']}（{info['role']}）" if info['role'] else info['name']
                shortages[key] = g['shortage']
        if g.get('run_violation'):
            run_violations[g['name']] = max_consecutive_run(g['days'])

    # パートスタッフの連続勤務チェック（固定パターンは変更しない。検知のみ）
    for info in rows_info:
        if info['is_full_time']:
            continue
        run = max_consecutive_run(set(info['pattern'].keys()))
        if run > MAX_CONSECUTIVE_WORK_DAYS:
            run_violations[info['name']] = run

    return {
        'ws': ws,
        'all_patterns': all_patterns,
        'code_by_key': code_by_key,
        'role_by_key': role_by_key,
        'days_in_month': days_in_month,
        'shortages': shortages,
        'run_violations': run_violations,
        'unmatched': unmatched,
        'cleared_rows': cleared_rows,
        'daily_counts': daily_counts,
        'understaffed_days': understaffed_days,
        'overstaffed_days': overstaffed_days,
        'biz_days': biz_days,
        'groups': groups,
        'svc_table': svc_table,
        'total_roles': total_roles,
    }


# ---------------------------------------------------------------------------
# レビュー（完成後の自動チェック）
# ---------------------------------------------------------------------------

def review_unit(result, staff_master, year, month, requests, unit_label):
    """完成したシフトを、決めたルールに沿っているか独立にもう一度チェックする。"""
    lines = [f'--- {unit_label} レビュー結果 ---']
    ok = True

    # (1) 常勤の実質所定日数
    keitai_by_name = {s['name']: s.get('keitai', '') for s in staff_master}
    target_days, _ = full_time_target_days(year, month)
    seen_names = set()
    for key, pattern in result['all_patterns'].items():
        name = key.split('（')[0]
        if keitai_by_name.get(name) not in ('A', 'B') or name in seen_names:
            continue
        seen_names.add(name)
        pl = len(paid_leave_days(name, requests, year, month))
        eff = max(0, target_days - pl)
        # pattern には有休の日（PAID_LEAVE_CODE）も含まれるようになったため、
        # 「実質所定日数」との比較は実際に稼働した日数（有休を除く）で行う。
        actual = sum(1 for code in pattern.values() if code != PAID_LEAVE_CODE)
        if actual != eff and key not in result['shortages']:
            ok = False
            lines.append(f'  ✗ {key}: 実質所定{eff}日に対し実績{actual}日（一致しません）')
    if result['shortages']:
        ok = False
        for key, s in result['shortages'].items():
            lines.append(f'  ✗ {key}: 候補日が足りず{s}日不足（要人手確認）')
    if ok:
        lines.append('  ✓ 常勤スタッフの実質所定日数: 全員一致')

    # (2) Total人数 4〜6人
    if result['understaffed_days']:
        ok = False
        days_str = '、'.join(f'{d}日' for d in result['understaffed_days'])
        lines.append(f'  ✗ Total人数が4人未満の営業日: {days_str}'
                      '（勤務可能希望があれば埋まります。なければ人員確保の相談が必要）')
    else:
        lines.append('  ✓ Total人数: 営業日はすべて4人以上')
    if result['overstaffed_days']:
        days_str = '、'.join(f'{d}日' for d in result['overstaffed_days'])
        lines.append(f'  ・参考: Total人数が6人を超える営業日（出力Excel上は赤字表示）: {days_str}'
                      '（自動では減らしていません。必要なら手動調整してください）')

    # (3) 連続勤務6日以上
    if result['run_violations']:
        ok = False
        for name, run in result['run_violations'].items():
            lines.append(f'  ✗ {name}: 連続勤務{run}日（6日以上）')
    else:
        lines.append('  ✓ 連続勤務: 6日以上の連続勤務なし')

    # (4) 兼務者（同一人物・複数行）の勤務日一致
    mismatch = False
    for g in result['groups']:
        if len(g['infos']) < 2:
            continue
        day_sets = [frozenset(info['pattern'].keys()) for info in g['infos']]
        if len(set(day_sets)) > 1:
            mismatch = True
            ok = False
            lines.append(f"  ✗ {g['name']}: 兼務している複数の行で勤務日が一致していません")
    if not mismatch and any(len(g['infos']) >= 2 for g in result['groups']):
        lines.append('  ✓ 兼務者の勤務日: 一致しています')

    # (5) staff_master.csvに存在しない退職者などの行を自動で氏名欄クリア
    if result.get('cleared_rows'):
        for r, name, role in result['cleared_rows']:
            lines.append(f'  ・自動処理: シート{r}行目（{name}／{role}）はstaff_master.csvに'
                         '見つからなかったため、氏名・職種・資格欄を自動で空にしました'
                         '（複製元の実績データは変更していません）')

    # (6) staff_master に一致しない行
    matched_names = {key.split('（')[0] for key in result['all_patterns']}
    real_gaps = []
    known_dupes = []
    for r, name, role in result['unmatched']:
        if name in matched_names:
            # 同じ人が他の行で正しくマッチしている（様式シート側の未使用の
            # 重複スロット、想定内）
            known_dupes.append((r, name, role))
        else:
            real_gaps.append((r, name, role))
    if known_dupes:
        for r, name, role in known_dupes:
            lines.append(f'  ・参考: シート{r}行目（{name}／{role}）は未使用の重複行です（{name}の勤務は他の行に記載・想定通り）')
    if real_gaps:
        ok = False
        for r, name, role in real_gaps:
            lines.append(f'  ✗ シート{r}行目（{name}／{role}）がstaff_master.csvに見つかりません。'
                         '新規スタッフの登録漏れ、または退職者の行が残っている可能性があります。'
                         '手作業で氏名欄を確認してください')
    if not result['unmatched']:
        lines.append('  ✓ 全スタッフ行がstaff_master.csvと正しく対応')

    lines.append('  === 総合判定: ' + ('問題なし' if ok else '要確認あり（上記✗を参照）') + ' ===')
    return '\n'.join(lines), ok


# ---------------------------------------------------------------------------
# サマリーレポート
# ---------------------------------------------------------------------------

def build_report(result, staff_master, year, month, requests):
    keitai_by_name = {s['name']: s.get('keitai', '') for s in staff_master}
    target_days, _ = full_time_target_days(year, month)
    all_patterns = result['all_patterns']
    shortages = result['shortages']

    lines = []
    lines.append('=== スタッフ別 稼働日数（この案。常勤はステップ3の自動調整済み） ===')
    for key, pattern in all_patterns.items():
        name = key.split('（')[0]
        is_ft = keitai_by_name.get(name, '') in ('A', 'B')
        note = ''
        # pattern には有休の日（PAID_LEAVE_CODE）も含まれる。稼働日数としては
        # 有休を除いた実際の勤務日数で数える（ユーザー確定ルール、2026-10：
        # 有休は月合計時間数には含めるが、稼働日数としては別枠のまま）。
        worked_days = sum(1 for code in pattern.values() if code != PAID_LEAVE_CODE)
        if is_ft:
            pl_count = len(paid_leave_days(name, requests, year, month))
            target_effective = max(0, target_days - pl_count)
            diff = worked_days - target_effective
            pl_note = f'（うち有休{pl_count}日を除く実質所定{target_effective}日）' if pl_count else ''
            if key in shortages:
                note = (f'  ※実質所定{target_effective}日{pl_note}に対し{shortages[key]}日不足 '
                        f'← 候補日が足りません。要確認（ステップ3）')
            elif diff != 0:
                note = f'  ※実質所定{target_effective}日{pl_note}との差: {diff:+d}日 ← 要確認'
            else:
                note = f'  （実質所定{target_effective}日{pl_note}と一致・自動調整済み）'
        lines.append(f'  {key}: {worked_days}日{note}')

    lines.append('')
    lines.append('=== 日別・職種別 配置人数（この案。Total＝様式のTotal行と同じ集計方法） ===')
    lines.append('※ 事務作業用など、サービス提供時間に掛かっていないシフト記号（例: 管理者のb）は'
                 'Totalにカウントされません（様式の数式と同じ扱い）。')
    svc_table = result.get('svc_table', {})
    role_by_key = result.get('role_by_key', {})
    total_roles = result.get('total_roles', set())
    for day in range(1, result['days_in_month'] + 1):
        counts = {}
        for key, pattern in all_patterns.items():
            # その日実際にシートへ書き込まれたシフト記号（有休の日は'有'）で判定する。
            # 以前は氏名ごとの既定コード（code_by_key）を使っていたが、有休の日は
            # 既定コードと異なる記号（有）になるため、必ずpattern[day]を使う。
            if day in pattern and counts_toward_total(pattern[day], role_by_key.get(key),
                                                        svc_table, total_roles):
                role = key.split('（')[-1].rstrip('）') if '（' in key else '?'
                counts[role] = counts.get(role, 0) + 1
        summary = ', '.join(f'{k}:{v}' for k, v in counts.items())
        total = result['daily_counts'].get(day, 0)
        if day in result['understaffed_days']:
            flag = '  ⚠4人未満'
        elif day in result['overstaffed_days']:
            flag = '  ⚠6人超'
        else:
            flag = ''
        biz = '' if day in result['biz_days'] else '（休業日）'
        lines.append(f'  {day:2d}日{biz}: Total={total} [{summary}]{flag}')
    return '\n'.join(lines)


# ---------------------------------------------------------------------------
# main
# ---------------------------------------------------------------------------

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--base', required=True, help='ベースとなる横浜市様式Excel（実績のあるファイル）')
    ap.add_argument('--source-am', required=True, help='コピー元となる午前シート名 例: 9月1')
    ap.add_argument('--source-pm', required=True, help='コピー元となる午後シート名 例: 9月2')
    ap.add_argument('--year', type=int, required=True)
    ap.add_argument('--month', type=int, required=True)
    ap.add_argument('--staff', required=True, help='staff_master.csv')
    ap.add_argument('--requests', default=None, help='希望・休み希望・勤務可能 CSV（任意）')
    ap.add_argument('--out', required=True)
    ap.add_argument('--shift-codes', default='shift_codes.csv',
                     help='シフト記号→サービス提供時間内の勤務時間数の対応表（Total集計に使用）')
    args = ap.parse_args()

    staff_master = load_staff_master(args.staff)
    requests = load_requests(args.requests)
    svc_table = load_service_hours_table(args.shift_codes)

    wb = openpyxl.load_workbook(args.base)

    new_am = f'{args.month}月1'
    new_pm = f'{args.month}月2'

    result_am = process_unit(wb, args.source_am, new_am, 'AM', staff_master, requests, args.year, args.month,
                              svc_table)
    result_pm = process_unit(wb, args.source_pm, new_pm, 'PM', staff_master, requests, args.year, args.month,
                              svc_table)

    # 出力ファイルには、その月のシフト案（午前・午後の2タブ）だけを残す。
    # ベースファイル（シフト2026.xlsx）に含まれる記入方法・記載例・過去月の
    # シートなどは複製元として読み込んだだけなので、出力からは削除する
    # （ユーザー確定ルール、2026-10）。
    # ただし「シフト記号表（勤務時間帯）」と「プルダウン・リスト」は削除しない：
    # 午前・午後シート内の勤務時間数・サービス提供時間内の勤務時間数のVLOOKUP式や、
    # 職種・資格のプルダウン（データの入力規則）がこの2シートを参照しているため、
    # 消すとTotalなどの自動計算が0になったり、開いたときに警告が出たりする。
    KEEP_ALWAYS = {'シフト記号表（勤務時間帯）', 'プルダウン・リスト'}
    for sheet_name in list(wb.sheetnames):
        if sheet_name in (new_am, new_pm) or sheet_name in KEEP_ALWAYS:
            continue
        del wb[sheet_name]

    Path(args.out).parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out)

    report_am = build_report(result_am, staff_master, args.year, args.month, requests)
    report_pm = build_report(result_pm, staff_master, args.year, args.month, requests)
    review_am, ok_am = review_unit(result_am, staff_master, args.year, args.month, requests, new_am + '（午前）')
    review_pm, ok_pm = review_unit(result_pm, staff_master, args.year, args.month, requests, new_pm + '（午後）')

    manual_staff = [
        s['name'] for s in staff_master
        if s.get('keitai') in ('C', 'D')
        and not s.get('fixed_weekdays_am', '').strip()
        and not s.get('fixed_weekdays_pm', '').strip()
    ]
    manual_staff = sorted(set(manual_staff))

    target_days, off_days = full_time_target_days(args.year, args.month)

    report_path = str(Path(args.out).with_suffix('')) + '_report.txt'
    with open(report_path, 'w', encoding='utf-8') as f:
        f.write(f'{args.year}年{args.month}月: 所定勤務日数（常勤） = {target_days}日 '
                f'（土日祝日 {off_days}日を除く）\n\n')
        if manual_staff:
            f.write('※ 以下のスタッフは不定期勤務のため自動生成の対象外です。'
                     '手作業でシフト記号を入力してください:\n')
            f.write('  ' + '、'.join(manual_staff) + '\n\n')
        f.write(f'--- {new_am}（午前） ---\n{report_am}\n\n')
        f.write(f'--- {new_pm}（午後） ---\n{report_pm}\n\n')
        f.write('\n=== レビュー（完成後の自動チェック） ===\n\n')
        f.write(review_am + '\n\n')
        f.write(review_pm + '\n')

    print(f'出力: {args.out}')
    print(f'レポート: {report_path}')
    print('')
    if manual_staff:
        print('不定期勤務のため未反映（手作業で入力）:', '、'.join(manual_staff))
    print('レビュー結果:', '午前=' + ('OK' if ok_am else '要確認'), '/', '午後=' + ('OK' if ok_pm else '要確認'))
    print('詳細はレポートファイルのレビュー欄を確認してください。')


if __name__ == '__main__':
    main()
